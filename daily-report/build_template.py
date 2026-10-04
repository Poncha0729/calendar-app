"""日報・日次決算テンプレート（年商1000億への10年計画）を生成する。

使い方:
    pip install openpyxl
    python3 build_template.py            # daily_report_template.xlsx を出力

出力した .xlsx は Google ドライブにアップロードして Google スプレッドシートに
変換して使う（Excel でもそのまま開ける）。
"""

import html
import os
import re
import zipfile
from datetime import date
from pathlib import Path

# openpyxl は lxml があると出力の書式が変わる（日本語が数値文字参照になる等）。常に同じ形で出力する。
os.environ.setdefault("OPENPYXL_LXML", "False")

from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.formula.translate import Translator  # noqa: E402
from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

OUT = Path(__file__).with_name("daily_report_template.xlsx")

# ---------------------------------------------------------------- シート名
S_HOWTO = "使い方"
S_FORM = "本日の日報"
S_DAYS = "月間日報"
S_HABIT = "毎日やること"
S_EXP = "経費ログ"
S_DEAL = "案件パイプライン"
S_MONTH = "月次集計"
S_PLAN = "10年計画"
S_LOG = "日次ログ"
S_SUBMIT = "提出データ"
S_SET = "設定"


def q(sheet):
    """数式用にシート名を引用符で囲む。"""
    return f"'{sheet}'"


# ---------------------------------------------------------------- スタイル
FONT = "Arial"
NAVY = "1F3864"
C_INPUT = "FFF2CC"  # 入力セル（Apps Script もこの色で入力欄を判定する）
C_HEAD = "D9E1F2"
C_SUB = "EDF1F9"
C_EXAMPLE = "F2F2F2"

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

FMT_YEN = '#,##0;▲#,##0;"-"'
FMT_CNT = '#,##0;▲#,##0;"-"'
FMT_DEC = '#,##0.0;▲#,##0.0;"-"'
FMT_PCT = '0.0%;▲0.0%;"-"'
FMT_DATE = "yyyy/mm/dd"
FMT_MONTH = "yyyy/mm"
FMT_DT = "yyyy/mm/dd hh:mm"

# 1000億への日課: (分類, 毎日やること, 短縮名, 具体的な行動, 目安, 1000億につながる理由)
HABITS = [
    ("数字", "目標から逆算して今日の数字を決める", "逆算・目標確認",
     "10年計画→今年→今月の目標と実績の差を確認し、今日の売上と行動件数の目標を決める",
     "朝5分",
     "10年目の目標を365で割ると、1日に必要な売上がわかる（年1000億円なら1日約2.7億円）。毎日差を確認しないと、差は月末まで気づかれないまま広がる"),
    ("数字", "日次決算を締める", "日次決算",
     "今日の売上・原価・経費・現預金残高をその日のうちに確定し、日報で提出する",
     "夜10分",
     "急成長中の会社は黒字でも資金が尽きることがある。利益と現金を毎日つかんでいれば、手を打つのが早くなる"),
    ("売上", "最重要タスクを午前中に終わらせる", "最重要タスク",
     "前回の日報で決めたTop3のうち、売上に最も効く1つを最初に片付ける",
     "午前中",
     "急ぎではないが重要な仕事は、後回しにするといつまでも終わらない。売上を動かす仕事を1日の最初に固定する"),
    ("売上", "新規アプローチを目標件数やる", "新規アプローチ",
     "架電・DM・訪問・紹介依頼などで、見込み客との新しい接点をつくる",
     "本日の日報の目標件数",
     "受注は商談から、商談は接点から生まれる。入口の件数が足りなければ、成約率が高くても売上は伸びない"),
    ("売上", "既存顧客と話す", "既存顧客",
     "最低1社と話し、課題・不満・追加の要望を聞く。満足している顧客には紹介を頼む",
     "1社以上",
     "既存顧客からの追加受注と紹介は、新規開拓より獲得コストが低い。解約の兆しにも早く気づける"),
    ("売上", "商談を前に進める", "商談前進",
     "案件パイプラインを更新し、すべての案件に次のアクションと期日を入れる。問い合わせや依頼には24時間以内に返す",
     "全案件",
     "次の一手が決まっていない案件は失注しやすい。止まっている案件をゼロにする"),
    ("仕組み", "仕事を1つ仕組み化する", "仕組み化",
     "今日やった仕事から1つ選び、マニュアル・テンプレート・チェックリストにして共有する",
     "1つ",
     "年1000億は自分1人では届かない。誰がやっても同じ成果が出る手順があれば、新しく入った人が早く戦力になる"),
    ("仕組み", "仕事を1つ手放す", "手放す",
     "自分がやらなくてよい仕事を1つ、人に任せる・自動化する・やめる",
     "1つ",
     "社長の時間は増やせない。手放して空いた時間を、売上・採用・提携に使う"),
    ("組織", "採用の接点を1つつくる", "採用",
     "候補者と話す、スカウトを送る、知人に紹介を頼む、採用ページを直す、のいずれかをやる",
     "1件",
     "1人あたり売上が1億円でも、年1000億には1,000人が必要になる。採用は毎日進めないと間に合わない"),
    ("組織", "メンバー1人と向き合う", "メンバー",
     "1on1・承認・感謝・フィードバックのいずれかを1人に行う",
     "1人",
     "採用しても定着しなければ組織は大きくならない。育成と定着が採用の成果を決める"),
    ("商品", "顧客の声から改善を1つ実行する", "改善",
     "商品・サービス・業務のどこかを、顧客の声をもとに1つ直す",
     "1つ",
     "毎日1つの改善は1年で約365個になる。競合との差は日々の改善の積み重ねでつく"),
    ("発信", "1つ発信する", "発信",
     "SNS・ブログ・導入事例・採用広報などで1つ発信する",
     "1本",
     "発信は問い合わせ・採用応募・提携の相談を生む。知られていない会社には相談が来ない"),
    ("拡大", "提携・資金・人脈に1手打つ", "提携・資金",
     "金融機関・投資家・提携先・業界のキーパーソン・M&A候補のいずれかに連絡するか会う",
     "1件",
     "年1000億には自社の営業だけでなく、資金調達・提携・M&Aによる成長が必要になる。関係づくりは必要になる前に始める"),
    ("学び", "30分学ぶ", "学習",
     "経営・財務・業界の動向や、先に年商1000億に届いた会社の事例を学び、1つを自社に当てはめる",
     "30分",
     "会社の規模が10倍になるたびに、経営者に必要な知識と判断が変わる"),
    ("自己管理", "体調を整える", "体調",
     "睡眠を7時間とり、運動する。体調を5段階で記録する",
     "睡眠7時間・運動20分",
     "10年間、判断の質を落とさずに働き続けるための土台になる"),
    ("振り返り", "日報を書いて提出する", "日報提出",
     "今日の結果を記録し、KPT（続けること・課題・次に試すこと）と明日のTop3を書いて提出する",
     "夜15分",
     "何が効いたかを毎日記録すると、成果が出たやり方を再現できる。記録がなければ振り返りは記憶頼みになる"),
]
N_HABIT = len(HABITS)

WEEKLY = [
    ("毎週", "月次集計で日課の実行率を見て、続かない日課の原因を1つ決めて手を打つ"),
    ("毎週", "案件パイプラインを全件見直し、止まっている案件の次の一手を決める"),
    ("毎週", "翌週の最重要テーマと売上目標を決める"),
    ("毎月", "月次集計で売上・粗利・営業利益を目標と比べ、差の原因を書き出す"),
    ("毎月", "設定シートの平均受注単価・受注率・商談化率を実績に合わせて更新する"),
    ("毎月", "採用・組織の計画を、10年計画の必要な営業担当者数と比べて見直す"),
]

THEMES = [
    "再現できる売り方を確立する：顧客・商品・売り方を絞り込み、営業手順を文書にする",
    "売り方を人に移す：営業マニュアルを整え、最初の営業チームを採用する",
    "組織で売る：マネジャーを育て、KPIで管理する体制をつくる",
    "販路を増やす：マーケティング・紹介・パートナー経由の売上を増やす",
    "資金と管理体制：資金調達と管理部門を強化し、上場の準備を検討する",
    "2つ目の事業：新商品・新事業・新エリアに展開する",
    "M&Aと提携：買収や提携で顧客・人材・商品を増やす",
    "経営チーム：事業ごとに責任者を置き、自分は方針と資本政策に集中する",
    "複数事業で伸ばす：事業の組み合わせで成長を続け、1つの市場の不振に左右されにくくする",
    "年商1000億円：次の10年の方針を決める",
]

# 勘定科目。「原価：」で始まる科目を売上原価、それ以外を販管費として集計する。
COGS_PREFIX = "原価："
ACCOUNTS = [
    "原価：仕入高", "原価：外注費", "販管費：広告宣伝費", "販管費：旅費交通費", "販管費：交際費",
    "販管費：会議費", "販管費：通信費", "販管費：消耗品費", "販管費：地代家賃", "販管費：水道光熱費",
    "販管費：支払手数料", "販管費：新聞図書費", "販管費：採用教育費", "販管費：給与・人件費",
    "販管費：福利厚生費", "販管費：保険料", "販管費：租税公課", "販管費：雑費",
]
PAYMENTS = ["現金", "クレジットカード", "銀行振込", "口座引落", "電子マネー・QR", "その他"]
PHASES = [("リード", 0.05), ("アポ獲得", 0.10), ("商談中", 0.25), ("提案済", 0.50),
          ("交渉中", 0.75), ("受注", 1.00), ("失注", 0.0), ("保留", 0.0)]


# ---------------------------------------------------------------- 書式ヘルパー
def font(bold=False, size=10, color="000000", italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def put(ws, ref, value=None, *, bold=False, size=10, color="000000", bg=None,
        fmt=None, align=None, wrap=False, border=False, italic=False, valign="center"):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font(bold, size, color, italic)
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
    if border:
        c.border = BORDER
    return c


def inp(ws, ref, value=None, fmt=None, align=None, wrap=False):
    """入力セル（薄い黄色・青文字）。"""
    return put(ws, ref, value, color="0000FF", bg=C_INPUT, fmt=fmt, align=align,
               wrap=wrap, border=True)


def calc(ws, ref, formula, fmt=None, link=False, align=None, wrap=False, bold=False):
    """計算セル（黒文字、他シート参照は緑文字）。"""
    return put(ws, ref, formula, color="008000" if link else "000000", fmt=fmt,
               align=align, wrap=wrap, border=True, bold=bold)


def title(ws, ref, text, merge_to=None):
    put(ws, ref, text, bold=True, size=14, color=NAVY)
    if merge_to:
        ws.merge_cells(f"{ref}:{merge_to}")


def note(ws, ref, text, merge_to=None, height=None):
    put(ws, ref, text, size=9, color="595959", wrap=True, valign="top")
    if merge_to:
        ws.merge_cells(f"{ref}:{merge_to}")
    if height:
        ws.row_dimensions[ws[ref].row].height = height


def section(ws, row, text, first="A", last="F"):
    put(ws, f"{first}{row}", text, bold=True, size=11, color="FFFFFF", bg=NAVY)
    ws.merge_cells(f"{first}{row}:{last}{row}")
    ws.row_dimensions[row].height = 20


def header(ws, row, labels, start_col=1, height=None):
    for i, label in enumerate(labels):
        ref = f"{get_column_letter(start_col + i)}{row}"
        put(ws, ref, label, bold=True, size=10, color=NAVY, bg=C_HEAD, align="center",
            wrap=True, border=True)
    if height:
        ws.row_dimensions[row].height = height


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def list_validation(ws, formula, ranges, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "リストから選んでください"
    dv.errorTitle = "入力エラー"
    if prompt:
        dv.prompt = prompt
    ws.add_data_validation(dv)
    for r in ranges:
        dv.add(r)


def nonneg_validation(ws, ranges):
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                        allow_blank=True)
    dv.error = "0以上の数値を入力してください"
    dv.errorTitle = "入力エラー"
    ws.add_data_validation(dv)
    for r in ranges:
        dv.add(r)


def status_colors(ws, rng):
    for sym, color in (("○", "D9EAD3"), ("△", "FCE5CD"), ("×", "F4CCCC")):
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=[f'"{sym}"'], fill=fill(color)))


# ---------------------------------------------------------------- 日次ログの列
# (キー, 見出し, 提出データ2行目の数式 or None(日課), 数値書式, 列幅)
F = {}  # 本日の日報の入力・計算セル（下で決まる）


def log_columns():
    fs = q(S_FORM)
    cols = [
        ("date", "日付", f"={fs}!{F['date']}", FMT_DATE, 12),
        ("dow", "曜日", f"={fs}!{F['dow']}", None, 5),
        ("year", "計画年目", f"={fs}!{F['year']}", "0", 6),
        ("health", "体調", f'=IF({fs}!{F["health"]}="","",{fs}!{F["health"]})', "0", 5),
        ("theme", "今日のテーマ", f"={fs}!{F['theme']}&\"\"", None, 22),
        ("t1r", "Top3①結果", f"={fs}!{F['top_res'][0]}&\"\"", None, 7),
        ("t2r", "Top3②結果", f"={fs}!{F['top_res'][1]}&\"\"", None, 7),
        ("t3r", "Top3③結果", f"={fs}!{F['top_res'][2]}&\"\"", None, 7),
        ("toprate", "Top3実行率", f"={fs}!{F['top_rate']}", FMT_PCT, 8),
        ("sales", "売上高", f"={fs}!{F['sales']}", FMT_YEN, 11),
        ("cogs", "売上原価", f"={fs}!{F['cogs']}", FMT_YEN, 11),
        ("gross", "粗利", f"={fs}!{F['gross']}", FMT_YEN, 11),
        ("sga", "販管費", f"={fs}!{F['sga']}", FMT_YEN, 11),
        ("op", "営業利益", f"={fs}!{F['op']}", FMT_YEN, 11),
        ("cash_in", "入金額", f"={fs}!{F['cash_in']}", FMT_YEN, 11),
        ("cash_out", "出金額", f"={fs}!{F['cash_out']}", FMT_YEN, 11),
        ("cash", "現預金残高", f'=IF({fs}!{F["cash"]}="","",{fs}!{F["cash"]})', FMT_YEN, 12),
        ("approach", "新規アプローチ", f"={fs}!{F['approach']}", FMT_CNT, 8),
        ("appo", "アポ獲得", f"={fs}!{F['appo']}", FMT_CNT, 7),
        ("meeting", "商談", f"={fs}!{F['meeting']}", FMT_CNT, 7),
        ("proposal", "提案・見積", f"={fs}!{F['proposal']}", FMT_CNT, 7),
        ("won", "受注数", f"={fs}!{F['won']}", FMT_CNT, 7),
        ("won_amt", "受注金額", f"={fs}!{F['won_amt']}", FMT_YEN, 11),
        ("existing", "既存顧客接点", f"={fs}!{F['existing']}", FMT_CNT, 7),
        ("referral", "紹介獲得", f"={fs}!{F['referral']}", FMT_CNT, 7),
    ]
    for i in range(N_HABIT):
        cols.append((f"h{i + 1}", None, f"={fs}!{F['habit_res'][i]}&\"\"", None, 7))
    cols += [
        ("habit_rate", "日課達成率", f"={fs}!{F['habit_rate']}", FMT_PCT, 8),
        ("result", "今日の成果", f"={fs}!{F['result']}&\"\"", None, 28),
        ("keep", "Keep", f"={fs}!{F['keep']}&\"\"", None, 28),
        ("problem", "Problem", f"={fs}!{F['problem']}&\"\"", None, 28),
        ("try", "Try", f"={fs}!{F['try']}&\"\"", None, 28),
        ("learn", "学び・気づき", f"={fs}!{F['learn']}&\"\"", None, 28),
        ("n1", "明日のTop3①", f"={fs}!{F['next'][0]}&\"\"", None, 24),
        ("n2", "明日のTop3②", f"={fs}!{F['next'][1]}&\"\"", None, 24),
        ("n3", "明日のTop3③", f"={fs}!{F['next'][2]}&\"\"", None, 24),
        ("consult", "相談・決めること", f"={fs}!{F['consult']}&\"\"", None, 24),
        ("submitted", "提出日時", "=NOW()", FMT_DT, 16),
    ]
    return cols


LOG_GROUPS = [  # (最初のキー, 最後のキー, 帯の見出し)
    ("date", "theme", "基本情報"),
    ("t1r", "toprate", "今日のTop3の結果"),
    ("sales", "cash", "日次決算（円）"),
    ("approach", "referral", "営業実績"),
    ("h1", "habit_rate", "1000億への日課（○△×）"),
    ("result", "learn", "振り返り"),
    ("n1", "consult", "明日の計画"),
    ("submitted", "submitted", "提出"),
]


def build():
    wb = Workbook()
    ws_howto = wb.active
    ws_howto.title = S_HOWTO
    ws_form = wb.create_sheet(S_FORM)
    ws_days = wb.create_sheet(S_DAYS)
    ws_habit = wb.create_sheet(S_HABIT)
    ws_exp = wb.create_sheet(S_EXP)
    ws_deal = wb.create_sheet(S_DEAL)
    ws_month = wb.create_sheet(S_MONTH)
    ws_plan = wb.create_sheet(S_PLAN)
    ws_log = wb.create_sheet(S_LOG)
    ws_sub = wb.create_sheet(S_SUBMIT)
    ws_set = wb.create_sheet(S_SET)

    # ============================================================ 設定
    ws = ws_set
    title(ws, "A1", "設定（目標・前提条件・選択肢）", "D1")
    note(ws, "A2", "薄い黄色のセルを自分の数字に書き換える。ほかのシートの計算はすべてここを参照している。",
         "D2", 18)
    header(ws, 4, ["項目", "値", "単位", "メモ（出所）"])
    settings = [
        ("start", "計画開始日", date(2026, 10, 1), "日付", FMT_DATE,
         "仮の値：1年目の初日。変更すると10年計画と月次集計の期間が連動する"),
        ("t1", "1年目の売上目標", 2, "億円", FMT_DEC,
         "仮の値：現状の売上から見て現実的な1年目の目標に変更する"),
        ("t10", "10年目の売上目標", 1000, "億円", FMT_DEC,
         "出所：ご本人の目標（10年後に年商1000億円）"),
        ("price", "平均受注単価（想定）", 500000, "円", FMT_YEN,
         "仮の値：直近の受注実績の平均に変更する"),
        ("close", "商談からの受注率（想定）", 0.3, "％", FMT_PCT,
         "仮の値：直近の実績（受注数÷商談数）に変更する"),
        ("meet", "アプローチからの商談化率（想定）", 0.1, "％", FMT_PCT,
         "仮の値：直近の実績（商談数÷新規アプローチ数）に変更する"),
        ("cap", "営業1人が1日にできるアプローチ数", 50, "件", FMT_CNT,
         "仮の値：10年計画の「必要な営業担当者数」の計算に使う"),
        ("exist", "既存顧客との接点（1日の目標）", 1, "社", FMT_CNT,
         "毎日やることNo.5の目安"),
        ("name", "記入者", None, "", None, "自分の名前（本日の日報と提出メールに表示）"),
        ("mail", "提出先メールアドレス", None, "", None,
         "Apps Scriptで提出したときにメールを送る宛先。空欄なら送らない"),
        ("remind", "リマインドの時刻", 21, "時", "0",
         "Apps Scriptの「毎日のリマインドを設定」で、未提出を知らせる時刻（0〜23）"),
    ]
    SET = {}
    for i, (key, label, val, unit, fmt, memo) in enumerate(settings):
        r = 5 + i
        put(ws, f"A{r}", label, border=True)
        inp(ws, f"B{r}", val, fmt=fmt)
        put(ws, f"C{r}", unit, border=True, align="center")
        put(ws, f"D{r}", memo, size=9, color="595959", border=True, wrap=True)
        SET[key] = f"{q(S_SET)}!$B${r}"
    # 選択肢リスト
    header(ws, 4, ["勘定科目"], start_col=6)
    for i, acc in enumerate(ACCOUNTS):
        inp(ws, f"F{5 + i}", acc)
    for r in range(5 + len(ACCOUNTS), 31):  # 追加用の空き行
        inp(ws, f"F{r}")
    note(ws, "F31", f"「{COGS_PREFIX}」で始まる科目は売上原価、それ以外は販管費として集計する。"
         f"空いている行に科目を追加できる（例：{COGS_PREFIX}材料費）。", "F34")
    header(ws, 4, ["支払方法"], start_col=9)
    for i, p in enumerate(PAYMENTS):
        inp(ws, f"I{5 + i}", p)
    for r in range(5 + len(PAYMENTS), 13):
        inp(ws, f"I{r}")
    header(ws, 4, ["案件フェーズ", "確度"], start_col=11)
    for i, (ph, prob) in enumerate(PHASES):
        inp(ws, f"K{5 + i}", ph)
        inp(ws, f"L{5 + i}", prob, fmt="0%")
    note(ws, "K13", "確度は仮の値。フェーズ別の加重見込金額（案件パイプライン）に使う。", "L15")
    widths(ws, {"A": 34, "B": 16, "C": 6, "D": 58, "E": 3, "F": 24, "G": 3, "H": 3,
                "I": 18, "J": 3, "K": 14, "L": 8})
    ws.freeze_panes = "A5"

    ACC_NAMES = f"{q(S_SET)}!$F$5:$F$30"
    PAY_LIST = f"{q(S_SET)}!$I$5:$I$12"
    PHASE_LIST = f"{q(S_SET)}!$K$5:$K$12"
    PH_WON = f"{q(S_SET)}!$K$10"
    PH_LOST = f"{q(S_SET)}!$K$11"

    # ============================================================ 毎日やること
    ws = ws_habit
    title(ws, "A1", "年商1000億円（10年後）に向けて毎日やること", "G1")
    ws["A2"] = (f'="10年目の目標 年"&TEXT({SET["t10"]},"#,##0")&"億円は、1日あたり約"'
                f'&TEXT({SET["t10"]}/365,"#,##0.0")&"億円。毎日の行動を下の16項目に分け、「本日の日報」で○△×をつける。"')
    put(ws, "A2", size=10, color="000000", wrap=True)
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 20
    note(ws, "A3", "項目名・短縮名・目安は自分の事業に合わせて書き換えてよい（本日の日報・日次ログ・月次集計に自動で反映される）。",
         "G3", 16)
    header(ws, 4, ["No", "分類", "毎日やること", "短縮名（ログ用）", "具体的な行動", "目安",
                   "1000億につながる理由"], height=30)
    HABIT_ROW0 = 5
    for i, (cat, name, short, action, guide, reason) in enumerate(HABITS):
        r = HABIT_ROW0 + i
        put(ws, f"A{r}", i + 1, border=True, align="center", valign="top")
        put(ws, f"B{r}", cat, border=True, align="center", valign="top")
        inp(ws, f"C{r}", name, wrap=True)
        inp(ws, f"D{r}", short, wrap=True)
        inp(ws, f"E{r}", action, wrap=True)
        inp(ws, f"F{r}", guide, wrap=True)
        put(ws, f"G{r}", reason, border=True, wrap=True, valign="top")
        for col in "CDEF":
            ws[f"{col}{r}"].alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 58
    r = HABIT_ROW0 + N_HABIT + 1
    section(ws, r, "週に1回・月に1回やること（日報の振り返りを計画に戻す）", "A", "G")
    header(ws, r + 1, ["", "頻度", "やること"])
    ws.merge_cells(f"C{r + 1}:G{r + 1}")
    for i, (freq, text) in enumerate(WEEKLY):
        rr = r + 2 + i
        put(ws, f"A{rr}", i + 1, border=True, align="center")
        put(ws, f"B{rr}", freq, border=True, align="center")
        put(ws, f"C{rr}", text, border=True, wrap=True)
        ws.merge_cells(f"C{rr}:G{rr}")
    widths(ws, {"A": 5, "B": 9, "C": 26, "D": 13, "E": 44, "F": 14, "G": 52})
    ws.freeze_panes = "A5"

    def habit_ref(col, i, absolute=False):
        d = "$" if absolute else ""
        return f"{q(S_HABIT)}!{d}{col}{d}{HABIT_ROW0 + i}"

    # ============================================================ 本日の日報
    ws = ws_form
    title(ws, "A1", "本日の日報・日次決算", "F1")
    note(ws, "A2", "薄い黄色のセルに入力（青い文字）。白いセルは自動計算、緑の文字はほかのシートから参照。"
         "入力が終わったらメニュー「日報」→「本日の日報を提出」（スクリプトがない場合は「提出データ」シートの手順）。",
         "F2", 30)

    # 1. 基本情報
    section(ws, 4, "1. 基本情報")
    put(ws, "A5", "日付", border=True)
    inp(ws, "B5", date(2026, 10, 1), fmt=FMT_DATE, align="center")
    F["date"] = "B5"
    calc(ws, "C5", '=IF(B5="","",CHOOSE(WEEKDAY(B5),"日","月","火","水","木","金","土"))',
         align="center")
    F["dow"] = "C5"
    put(ws, "D5", "計画開始から", border=True)
    calc(ws, "E5", f'=IF(B5="","",B5-{SET["start"]}+1)', fmt='0"日目"', link=True, align="right")
    note(ws, "F5", "Ctrl + ;（セミコロン）で今日の日付を入力できる")

    put(ws, "A6", "計画年目", border=True)
    calc(ws, "B6", f'=MIN(10,IFERROR(DATEDIF({SET["start"]},B5,"Y")+1,1))', fmt='0"年目"',
         link=True, align="center")
    F["year"] = "B6"
    put(ws, "C6", None, border=True)
    put(ws, "D6", "前回の提出日", border=True)
    calc(ws, "E6", f"=_xlfn.MAXIFS({q(S_LOG)}!$A:$A,{q(S_LOG)}!$A:$A,\"<\"&$B$5)",
         fmt='yyyy/mm/dd;;"なし"', link=True, align="right")
    F["prev"] = "E6"
    note(ws, "F6", "日次ログにある、この日付より前の最新の日報")

    put(ws, "A7", "体調（1〜5）", border=True)
    inp(ws, "B7", fmt="0", align="center")
    F["health"] = "B7"
    put(ws, "C7", None, border=True)
    put(ws, "D7", "記入者", border=True)
    calc(ws, "E7", f'={SET["name"]}&""', link=True, align="right")
    note(ws, "F7", "5＝絶好調、3＝普通、1＝不調")

    put(ws, "A8", "今日のテーマ（一言）", border=True)
    inp(ws, "B8", wrap=True)
    ws.merge_cells("B8:F8")
    F["theme"] = "B8"

    # 2. 今日のTop3
    section(ws, 10, "2. 今日のTop3（前回の日報で決めたこと）")
    header(ws, 11, ["項目", "結果（○△×）", "内容（前回の日報から自動表示）", "", "", ""])
    ws.merge_cells("C11:F11")
    F["top_res"] = []
    next_cols = ["AV", "AW", "AX"]  # 日次ログの「明日のTop3」列（下で検証する）
    for i in range(3):
        r = 12 + i
        put(ws, f"A{r}", ["① 最重要", "②", "③"][i], border=True)
        inp(ws, f"B{r}", align="center")
        F["top_res"].append(f"B{r}")
        calc(ws, f"C{r}",
             f"=IFERROR(INDEX({q(S_LOG)}!${next_cols[i]}:${next_cols[i]},"
             f"MATCH($E$6,{q(S_LOG)}!$A:$A,0))&\"\",\"\")", link=True, wrap=True)
        ws.merge_cells(f"C{r}:F{r}")
        ws.row_dimensions[r].height = 30
    put(ws, "A15", "Top3の実行率", border=True)
    calc(ws, "B15", '=IF(COUNTIF(C12:C14,"?*")=0,0,COUNTIF(B12:B14,"○")/COUNTIF(C12:C14,"?*"))',
         fmt=FMT_PCT, align="center")
    F["top_rate"] = "B15"
    note(ws, "C15", "○の数÷Top3の数。前回の日報がない日は「-」", "F15")

    # 3. 目標と進捗
    section(ws, 17, "3. 売上目標と進捗（年商1000億からの逆算）")
    header(ws, 18, ["項目", "実績", "単位", "目標", "達成率", "メモ"])
    plan_d = f"{q(S_PLAN)}!$D$10:$D$19"
    month_start = "DATE(YEAR($B$5),MONTH($B$5),1)"
    year_start = f'EDATE({SET["start"]},12*($B$6-1))'
    log_a = f"{q(S_LOG)}!$A:$A"
    log_sales = f"{q(S_LOG)}!$J:$J"
    rows3 = [
        ("今日の売上", "=B38", "=D21/365", "目標＝今年の売上目標÷365"),
        ("今月の売上（本日まで）",
         f'=SUMIFS({log_sales},{log_a},">="&{month_start},{log_a},"<"&$B$5)+$B$38',
         "=D21/12", "目標＝今年の売上目標÷12"),
        ("今年の売上（本日まで）",
         f'=SUMIFS({log_sales},{log_a},">="&{year_start},{log_a},"<"&$B$5)+$B$38',
         f"=INDEX({plan_d},$B$6)*10^8", None),
    ]
    for i, (label, actual, target, memo) in enumerate(rows3):
        r = 19 + i
        put(ws, f"A{r}", label, border=True)
        calc(ws, f"B{r}", actual, fmt=FMT_YEN, link=(i > 0))
        put(ws, f"C{r}", "円", border=True, align="center")
        calc(ws, f"D{r}", target, fmt=FMT_YEN, link=(i == 2))
        calc(ws, f"E{r}", f"=IF(D{r}=0,0,B{r}/D{r})", fmt=FMT_PCT)
        if memo:
            note(ws, f"F{r}", memo)
    put(ws, "F21", '="計画"&$B$6&"年目の目標（10年計画シート）"', size=9, color="595959")
    put(ws, "A22", "今月の目標まであと", border=True)
    calc(ws, "B22", "=MAX(0,D20-B20)", fmt=FMT_YEN)
    put(ws, "C22", "円", border=True, align="center")
    note(ws, "F22", "今月の目標から今月の売上を引いた額")
    put(ws, "A23", "明日以降に必要な売上（1日あたり）", border=True)
    calc(ws, "B23", "=IF(EOMONTH($B$5,0)-$B$5<=0,0,B22/(EOMONTH($B$5,0)-$B$5))", fmt=FMT_YEN)
    put(ws, "C23", "円", border=True, align="center")
    note(ws, "F23", "残りの額を、月末までの残り日数で割った額")

    # 4. 営業実績
    section(ws, 25, "4. 営業実績（本日）")
    header(ws, 26, ["項目", "実績", "単位", "目標", "達成率", "メモ"])
    plan = lambda col: f"{q(S_PLAN)}!${col}$10:${col}$19"  # noqa: E731
    sales_rows = [
        ("approach", "新規アプローチ数", "件", f"=INDEX({plan('L')},$B$6)",
         "架電・DM・訪問・紹介依頼など、見込み客との新しい接点。目標は10年計画の逆算"),
        ("appo", "アポ獲得数", "件", None, None),
        ("meeting", "商談数（実施）", "件", f"=INDEX({plan('J')},$B$6)/365",
         "目標＝10年計画の年間必要商談数÷365"),
        ("proposal", "提案・見積の提出数", "件", None, None),
        ("won", "受注数", "件", f"=INDEX({plan('I')},$B$6)/365", "目標＝10年計画の年間必要受注件数÷365"),
        ("won_amt", "受注金額", "円", "=D19", "契約した金額（売上の計上日とずれてもよい）"),
        ("existing", "既存顧客との面談・連絡", "社", f'={SET["exist"]}', "毎日やることNo.5"),
        ("referral", "紹介でもらった見込み客", "件", None, None),
    ]
    for i, (key, label, unit, target, memo) in enumerate(sales_rows):
        r = 27 + i
        put(ws, f"A{r}", label, border=True)
        inp(ws, f"B{r}", fmt=FMT_YEN if unit == "円" else FMT_CNT)
        F[key] = f"B{r}"
        put(ws, f"C{r}", unit, border=True, align="center")
        if target:
            calc(ws, f"D{r}", target, fmt=FMT_YEN if unit == "円" else FMT_DEC,
                 link=(target.startswith(f"=INDEX") or "設定" in target))
            calc(ws, f"E{r}", f"=IF(D{r}=0,0,B{r}/D{r})", fmt=FMT_PCT)
        else:
            put(ws, f"D{r}", None, border=True)
            put(ws, f"E{r}", None, border=True)
        if memo:
            note(ws, f"F{r}", memo)

    # 5. 日次決算
    section(ws, 36, "5. 日次決算（本日）")
    header(ws, 37, ["項目", "本日", "単位", "今月累計", "", "メモ"])
    exp_a = f"{q(S_EXP)}!$A:$A"
    exp_b = f"{q(S_EXP)}!$B:$B"
    exp_d = f"{q(S_EXP)}!$D:$D"
    mtd = f'{exp_a},">="&{month_start},{exp_a},"<="&$B$5'

    def cogs(crit):
        return f'=SUMIFS({exp_d},{crit},{exp_b},"{COGS_PREFIX}*")'

    pl_rows = [
        ("sales", "売上高", None, "=B20", "本日計上した売上（税抜）"),
        ("cogs", "売上原価", cogs(f"{exp_a},$B$5"), cogs(mtd),
         f"経費ログのうち、科目が「{COGS_PREFIX}」で始まる費用の合計"),
        ("gross", "粗利", "=B38-B39", "=D38-D39", None),
        ("gross_rate", "粗利率", "=IF(B38=0,0,B40/B38)", "=IF(D38=0,0,D40/D38)", None),
        ("sga", "販管費", f"=SUMIFS({exp_d},{exp_a},$B$5)-B39", f"=SUMIFS({exp_d},{mtd})-D39",
         "経費ログのうち、原価以外の科目の合計"),
        ("op", "営業利益", "=B40-B42", "=D40-D42", None),
        ("op_rate", "営業利益率", "=IF(B38=0,0,B43/B38)", "=IF(D38=0,0,D43/D38)", None),
    ]
    for i, (key, label, today, month, memo) in enumerate(pl_rows):
        r = 38 + i
        pct = key.endswith("rate")
        put(ws, f"A{r}", label, border=True, bold=key in ("gross", "op"))
        if today is None:
            inp(ws, f"B{r}", fmt=FMT_YEN)
        else:
            calc(ws, f"B{r}", today, fmt=FMT_PCT if pct else FMT_YEN, link=key in ("cogs", "sga"),
                 bold=key in ("gross", "op"))
        F[key] = f"B{r}"
        put(ws, f"C{r}", "％" if pct else "円", border=True, align="center")
        calc(ws, f"D{r}", month, fmt=FMT_PCT if pct else FMT_YEN,
             link=key in ("sales", "cogs", "sga"), bold=key in ("gross", "op"))
        put(ws, f"E{r}", None, border=True)
        if memo:
            note(ws, f"F{r}", memo)
    put(ws, "A45", "経費の記録件数", border=True)
    calc(ws, "B45", f"=COUNTIFS({exp_a},$B$5)", fmt="0", link=True)
    put(ws, "C45", "件", border=True, align="center")
    note(ws, "F45", "費用は「経費ログ」に1件1行で入力。0件のときは記録漏れがないか確認")
    put(ws, "A46", "現預金（任意）", bold=True, color=NAVY, bg=C_SUB)
    ws.merge_cells("A46:F46")
    cash_rows = [
        ("prev_cash", "前回の現預金残高",
         f"=IFERROR(N(INDEX({q(S_LOG)}!$Q:$Q,MATCH($E$6,{log_a},0))),0)", "前回提出した日報の残高"),
        ("cash_in", "本日の入金額", None, None),
        ("cash_out", "本日の出金額", None, None),
        ("cash_calc", "計算上の残高", "=B47+B48-B49", "前回の残高＋入金−出金"),
        ("cash", "終業時の現預金残高（実際）", None, "銀行口座と手元現金の合計"),
        ("cash_diff", "差異（計算上−実際）", '=IF(OR($B$51="",$E$6=0),0,B50-B51)',
         "0でなければ、入出金の記録漏れがないか確認"),
    ]
    for i, (key, label, formula, memo) in enumerate(cash_rows):
        r = 47 + i
        put(ws, f"A{r}", label, border=True)
        if formula is None:
            inp(ws, f"B{r}", fmt=FMT_YEN)
        else:
            calc(ws, f"B{r}", formula, fmt=FMT_YEN, link=(key == "prev_cash"))
        F[key] = f"B{r}"
        put(ws, f"C{r}", "円", border=True, align="center")
        if memo:
            note(ws, f"F{r}", memo)
    ws.conditional_formatting.add(
        "B52", CellIsRule(operator="notEqual", formula=["0"], fill=fill("F4CCCC")))

    # 6. 日課チェック
    section(ws, 54, "6. 1000億への日課チェック（○＝できた、△＝一部、×＝できなかった）")
    header(ws, 55, ["日課", "結果（○△×）", "目安", "", "", "やったこと・メモ"])
    ws.merge_cells("C55:E55")
    F["habit_res"] = []
    for i in range(N_HABIT):
        r = 56 + i
        calc(ws, f"A{r}", f'={habit_ref("A", i)}&". "&{habit_ref("C", i)}', link=True, wrap=True)
        inp(ws, f"B{r}", align="center")
        F["habit_res"].append(f"B{r}")
        calc(ws, f"C{r}", f'={habit_ref("F", i)}&""', link=True, wrap=True)
        ws.merge_cells(f"C{r}:E{r}")
        inp(ws, f"F{r}", wrap=True)
        ws.row_dimensions[r].height = 28
    hr0, hr1 = 56, 56 + N_HABIT - 1
    put(ws, "A72", "日課の達成率", border=True, bold=True)
    calc(ws, "B72", f'=IF(COUNTIF(A{hr0}:A{hr1},"?*")=0,0,COUNTIF(B{hr0}:B{hr1},"○")/COUNTIF(A{hr0}:A{hr1},"?*"))',
         fmt=FMT_PCT, align="center", bold=True)
    F["habit_rate"] = "B72"
    calc(ws, "C72", f'="○ "&COUNTIF(B{hr0}:B{hr1},"○")&"　△ "&COUNTIF(B{hr0}:B{hr1},"△")'
                    f'&"　× "&COUNTIF(B{hr0}:B{hr1},"×")', align="center")
    ws.merge_cells("C72:E72")
    note(ws, "F72", "○の数÷日課の数")

    # 7. 振り返り
    section(ws, 74, "7. 振り返り（KPT）と明日の計画")
    kpt = [
        ("result", "今日いちばんの成果（数字で）", 36),
        ("keep", "Keep：うまくいった・続けること", 48),
        ("problem", "Problem：課題・うまくいかなかったこと", 48),
        ("try", "Try：次に試すこと", 48),
        ("learn", "学び・気づき", 36),
        ("n1", "明日のTop3 ①（最重要）", 24),
        ("n2", "明日のTop3 ②", 24),
        ("n3", "明日のTop3 ③", 24),
        ("consult", "相談したいこと・決めること", 36),
    ]
    for i, (key, label, h) in enumerate(kpt):
        r = 75 + i
        put(ws, f"A{r}", label, border=True, wrap=True)
        inp(ws, f"B{r}", wrap=True)
        ws[f"B{r}"].alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(f"B{r}:F{r}")
        ws.row_dimensions[r].height = h
        F[key] = f"B{r}"
    F["next"] = [F.pop("n1"), F.pop("n2"), F.pop("n3")]
    note(ws, "A85", "提出：メニュー「日報」→「本日の日報を提出」。日次ログに1行追加され、入力欄を空にして翌日の日付に進められる。"
         "スクリプトを入れていない場合は「提出データ」シートの2行目をコピーし、「日次ログ」の最終行の下に値のみ貼り付け（Ctrl+Shift+V）。"
         "提出した日報は「月間日報」に1日1列で並び、1か月分をまとめて見返せる。",
         "F85", 54)

    # 入力規則・条件付き書式
    list_validation(ws, '"1,2,3,4,5"', ["B7"])
    list_validation(ws, '"○,△,×"', ["B12:B14", f"B{hr0}:B{hr1}"])
    nonneg_validation(ws, ["B27:B34", "B38", "B48:B49", "B51"])
    status_colors(ws, f"B12:B14 B{hr0}:B{hr1}")
    for rng in ("E19:E21", "E27:E33"):
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=fill("D9EAD3")))
    ws.conditional_formatting.add(
        "B72", CellIsRule(operator="greaterThanOrEqual", formula=["0.8"], fill=fill("D9EAD3")))
    widths(ws, {"A": 34, "B": 16, "C": 7, "D": 16, "E": 11, "F": 48})
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = "F1C232"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ============================================================ 日次ログ / 提出データ
    cols = log_columns()
    L = {key: get_column_letter(i + 1) for i, (key, *_rest) in enumerate(cols)}
    assert [L["n1"], L["n2"], L["n3"]] == next_cols, L
    assert L["sales"] == "J" and L["cash"] == "Q", L
    last_col = get_column_letter(len(cols))

    def habit_header(i):
        return f'={habit_ref("A", i, True)}&" "&{habit_ref("D", i, True)}'

    ws = ws_log
    for first, last, label in LOG_GROUPS:
        a, b = L[first], L[last]
        put(ws, f"{a}1", label, bold=True, color="FFFFFF", bg=NAVY, align="center")
        if a != b:
            ws.merge_cells(f"{a}1:{b}1")
    example = {
        "date": "記入例", "dow": "木", "year": 1, "health": 4, "theme": "新規開拓に集中する日",
        "t1r": "○", "t2r": "○", "t3r": "△", "toprate": 2 / 3,
        "sales": 300000, "cogs": 90000, "gross": 210000, "sga": 25000, "op": 185000,
        "cash_in": 500000, "cash_out": 120000, "cash": 3200000,
        "approach": 35, "appo": 4, "meeting": 3, "proposal": 2, "won": 1, "won_amt": 500000,
        "existing": 2, "referral": 1,
        "result": "新規1件受注（50万円）", "keep": "午前中に架電をまとめたらアポ率が上がった",
        "problem": "見積の作成に2時間かかった", "try": "見積のテンプレートを作る",
        "learn": "決裁者が同席すると商談が早く進む", "n1": "A社に提案書を送る",
        "n2": "営業トークをマニュアルにする", "n3": "採用候補者と面談", "consult": "",
        "submitted": None,
    }
    pattern = "○○○○△○○×○△○○×○○○"
    for i, s in enumerate(pattern):
        example[f"h{i + 1}"] = s
    example["habit_rate"] = pattern.count("○") / N_HABIT
    for i, (key, label, formula, fmt, width) in enumerate(cols):
        col = L[key]
        if key.startswith("h") and key[1:].isdigit():
            header_value = habit_header(int(key[1:]) - 1)
        else:
            header_value = label
        for ws2 in (ws_log, ws_sub):
            put(ws2, f"{col}2" if ws2 is ws_log else f"{col}1", header_value, bold=True,
                color=NAVY, bg=C_HEAD, align="center", wrap=True, border=True)
            ws2.column_dimensions[col].width = width
        put(ws, f"{col}3", example[key], color="7F7F7F", italic=True, bg=C_EXAMPLE, fmt=fmt,
            border=True, align="center" if width <= 8 else None)
        if fmt:
            ws.column_dimensions[col].number_format = fmt
        calc(ws_sub, f"{col}2", formula, fmt=fmt, link=True)
    ws.row_dimensions[2].height = 45
    ws.freeze_panes = "B3"
    ws.sheet_properties.tabColor = "808080"

    ws = ws_sub
    ws.row_dimensions[1].height = 45
    note(ws, "A4", "このシートは自動作成（編集しない）。Apps Scriptの「本日の日報を提出」は、この2行目を日次ログに書き込む。"
         "スクリプトを使わない場合は、2行目（A2〜" + last_col + "2）をコピーし、「日次ログ」の最終行の下に"
         "値のみ貼り付け（Ctrl+Shift+V）する。日付が数字で表示されたら、日次ログのA列を選んで「表示形式」→「数字」→「日付」にする。", "P4", 44)
    ws.freeze_panes = "B1"
    ws.sheet_properties.tabColor = "808080"

    # 名前付き範囲（Apps Script が日付セルを探すため）
    wb.defined_names["REPORT_DATE"] = DefinedName(
        "REPORT_DATE", attr_text=f"{q(S_FORM)}!$B$5")

    # ============================================================ 月間日報
    # 提出した日報（日次ログ）を、本日の日報と同じ項目の並びで1日1列にし、1か月分を1シートに表示する。
    ws = ws_days
    D0 = 5  # 1日目の列（E列）。E〜AI列に1日〜31日を並べる
    day_cols = [get_column_letter(D0 + k) for k in range(31)]
    c0, c1 = day_cols[0], day_cols[-1]
    PCT_DAY = '0%;▲0%;"-"'

    def span(r):
        return f"{c0}{r}:{c1}{r}"

    def pull(key, col):
        src = f"{q(S_LOG)}!${L[key]}:${L[key]}"
        return f'=IF({col}$5="","",IF(INDEX({src},{col}$5)="","",INDEX({src},{col}$5)))'

    title(ws, "A1", "月間日報（1か月分の日報）", "D1")
    note(ws, f"{c0}1", "「本日の日報」で提出した日報が、1日1列で自動で並ぶ（このシートには入力しない）。"
         "提出していない日は空欄。月計・月目標・達成率は、表示している月の分。", f"{c1}1")
    put(ws, "A2", "表示する月", border=True, bold=True)
    inp(ws, "B2", fmt=FMT_DATE, align="center")
    dv = DataValidation(type="date", operator="greaterThan", formula1="1", allow_blank=True)
    dv.error = "日付を入力してください（例：2026/10/1）"
    dv.errorTitle = "入力エラー"
    ws.add_data_validation(dv)
    dv.add("B2")
    form_date = f"{q(S_FORM)}!$B$5"
    calc(ws, "C2", f"=IF(ISNUMBER($B$2),DATE(YEAR($B$2),MONTH($B$2),1),"
                   f"IF(ISNUMBER({form_date}),DATE(YEAR({form_date}),MONTH({form_date}),1),"
                   f"DATE(YEAR(TODAY()),MONTH(TODAY()),1)))",
         fmt='yyyy"年"m"月"', link=True, align="center", bold=True)
    calc(ws, "D2", f'=MIN(10,IFERROR(DATEDIF({SET["start"]},$C$2,"Y")+1,1))', fmt='0"年目"',
         link=True, align="center")
    note(ws, f"{c0}2", "B2が空欄なら「本日の日報」の日付の月を表示する。ほかの月はB2にその月の日付を入力（例：2026/10/1）。"
         "月ごとに残したいときは、このシートをコピーしてB2にその月の日付を入れる。", f"{c1}2", 28)

    header(ws, 3, ["項目", "月計", "月目標", "達成率"])
    header(ws, 4, ["曜日", "", "", ""])
    put(ws, "A5", "日報の提出（○＝提出済み）", border=True, bold=True)
    calc(ws, "B5", f"=COUNT({span(5)})", fmt='0"日"', align="center", bold=True)
    put(ws, "C5", None, border=True)
    put(ws, "D5", None, border=True)
    for k, col in enumerate(day_cols):
        prev = get_column_letter(D0 + k - 1)
        day = "=$C$2" if k == 0 else f'=IF({prev}3="","",IF(MONTH({prev}3+1)<>MONTH({prev}3),"",{prev}3+1))'
        put(ws, f"{col}3", day, bold=True, color=NAVY, bg=C_HEAD, align="center", border=True, fmt="m/d")
        put(ws, f"{col}4", f'=IF({col}3="","",CHOOSE(WEEKDAY({col}3),"日","月","火","水","木","金","土"))',
            bold=True, color=NAVY, bg=C_HEAD, align="center", border=True)
        # 日次ログの行番号（提出済みなら○と表示する）
        calc(ws, f"{col}5", f'=IF({col}3="","",IFERROR(MATCH({col}3,{log_a},0),""))', fmt='"○";;',
             link=True, align="center")

    # (種類, 日次ログのキー, 見出し, …)
    #   val: 数値（書式, 月計の種類, 月目標）  sym: ○△×（達成率の分母）  ratio: 率（分子キー, 分母キー）  text: 文章
    log_cash = f"{q(S_LOG)}!${L['cash']}:${L['cash']}"
    spec = [
        ("sec", "1. 基本情報"),
        ("val", "health", "体調（1〜5）", "0", "avg", None),
        ("text", "theme", "今日のテーマ"),
        ("sec", "2. 今日のTop3の結果（○△×）"),
        ("sym", "t1r", "Top3 ①（最重要）", "marked"),
        ("sym", "t2r", "Top3 ②", "marked"),
        ("sym", "t3r", "Top3 ③", "marked"),
        ("val", "toprate", "Top3の実行率", PCT_DAY, "avg", None),
        ("sec", "3. 日次決算（円）"),
        ("val", "sales", "売上高", FMT_YEN, "sum", f"=INDEX({plan_d},$D$2)*10^8/12"),
        ("val", "cogs", "売上原価", FMT_YEN, "sum", None),
        ("val", "gross", "粗利", FMT_YEN, "sum", None),
        ("ratio", "gross_rate", "粗利率", "gross", "sales"),
        ("val", "sga", "販管費", FMT_YEN, "sum", None),
        ("val", "op", "営業利益", FMT_YEN, "sum", None),
        ("ratio", "op_rate", "営業利益率", "op", "sales"),
        ("val", "cash_in", "入金額", FMT_YEN, "sum", None),
        ("val", "cash_out", "出金額", FMT_YEN, "sum", None),
        ("val", "cash", "現預金残高（月計は月末時点）", FMT_YEN, "last", None),
        ("sec", "4. 営業実績"),
        ("val", "approach", "新規アプローチ数", FMT_CNT, "sum", f"=INDEX({plan('K')},$D$2)/12"),
        ("val", "appo", "アポ獲得数", FMT_CNT, "sum", None),
        ("val", "meeting", "商談数", FMT_CNT, "sum", f"=INDEX({plan('J')},$D$2)/12"),
        ("val", "proposal", "提案・見積の提出数", FMT_CNT, "sum", None),
        ("val", "won", "受注数", FMT_CNT, "sum", f"=INDEX({plan('I')},$D$2)/12"),
        ("val", "won_amt", "受注金額", FMT_YEN, "sum", "sales"),
        ("val", "existing", "既存顧客との面談・連絡", FMT_CNT, "sum", f'={SET["exist"]}*DAY(EOMONTH($C$2,0))'),
        ("val", "referral", "紹介でもらった見込み客", FMT_CNT, "sum", None),
        ("sec", "5. 1000億への日課（○△×）"),
        *[("sym", f"h{i + 1}", f'={habit_ref("A", i)}&" "&{habit_ref("D", i)}', "submitted")
          for i in range(N_HABIT)],
        ("val", "habit_rate", "日課の達成率", PCT_DAY, "avg", None),
        ("sec", "6. 振り返り（KPT）"),
        ("text", "result", "今日いちばんの成果"),
        ("text", "keep", "Keep：続けること"),
        ("text", "problem", "Problem：課題"),
        ("text", "try", "Try：次に試すこと"),
        ("text", "learn", "学び・気づき"),
        ("sec", "7. 明日の計画"),
        ("text", "n1", "明日のTop3 ①（最重要）"),
        ("text", "n2", "明日のTop3 ②"),
        ("text", "n3", "明日のTop3 ③"),
        ("text", "consult", "相談・決めること"),
    ]
    R = {}
    r = 6
    for item in spec:
        kind = item[0]
        if kind == "sec":
            section(ws, r, item[1], "A", c1)
            r += 1
            continue
        key, label = item[1], item[2]
        R[key] = r
        if label.startswith("="):
            calc(ws, f"A{r}", label, link=True)
        else:
            put(ws, f"A{r}", label, border=True)
        if kind == "ratio":
            num, den = R[item[3]], R[item[4]]
            for col in day_cols:
                calc(ws, f"{col}{r}", f'=IF(N({col}{den})=0,"",N({col}{num})/{col}{den})', fmt=PCT_DAY)
            calc(ws, f"B{r}", f'=IF(N(B{den})=0,"",B{num}/B{den})', fmt=FMT_PCT)
            put(ws, f"C{r}", None, border=True)
            put(ws, f"D{r}", None, border=True)
            r += 1
            continue
        fmt = item[3] if kind == "val" else None
        for col in day_cols:
            c = calc(ws, f"{col}{r}", pull(key, col), fmt=fmt, link=True,
                     align="center" if kind == "sym" else None)
            if kind == "text":
                c.alignment = Alignment(vertical="top", wrap_text=True)
        if kind == "val":
            summary, target = item[4], item[5]
            total = {
                "sum": f"=SUM({span(r)})",
                "avg": f'=IFERROR(AVERAGE({span(r)}),"")',
                "last": (f'=IFERROR(INDEX({log_cash},MATCH(_xlfn.MAXIFS({log_a},{log_a},">="&$C$2,'
                         f'{log_a},"<="&EOMONTH($C$2,0),{log_cash},"<>"),{log_a},0)),"")'),
            }[summary]
            calc(ws, f"B{r}", total, fmt=FMT_PCT if fmt == PCT_DAY else fmt, link=summary == "last")
            if target == "sales":
                target = f"=C{R['sales']}"
            if target:
                calc(ws, f"C{r}", target, fmt=FMT_YEN if fmt == FMT_YEN else FMT_DEC, link=True)
                calc(ws, f"D{r}", f'=IF(N(C{r})=0,"",B{r}/C{r})', fmt=FMT_PCT)
            else:
                put(ws, f"C{r}", None, border=True)
                put(ws, f"D{r}", None, border=True)
        elif kind == "sym":
            calc(ws, f"B{r}", f'=COUNTIF({span(r)},"○")', fmt='0"回"', align="center")
            put(ws, f"C{r}", None, border=True)
            base = f'COUNTIF({span(r)},"?*")' if item[3] == "marked" else "$B$5"
            calc(ws, f"D{r}", f'=IF({base}=0,"",B{r}/{base})', fmt=FMT_PCT)
        else:
            for col in "BCD":
                put(ws, f"{col}{r}", None, border=True, bg=C_EXAMPLE)
        r += 1
    note(ws, f"A{r + 1}", "このシートは「日次ログ」から自動で表示している（入力しない）。"
         "Top3の達成率は○の数÷結果を入れた日数、日課の達成率は○の数÷日報を提出した日数。"
         "売上原価・販管費は提出した時点の値（あとで経費ログを直した分は「月次集計」に反映される）。",
         f"{get_column_letter(D0 + 10)}{r + 1}", 42)

    status_colors(ws, f"{c0}{R['t1r']}:{c1}{R['t3r']} {c0}{R['h1']}:{c1}{R[f'h{N_HABIT}']}")
    for dow, color in (("土", "1155CC"), ("日", "C00000")):
        ws.conditional_formatting.add(
            f"{c0}3:{c1}4", FormulaRule(formula=[f'{c0}$4="{dow}"'], font=Font(color=color, bold=True)))
    d0, d1 = R["sales"], R["referral"]
    ws.conditional_formatting.add(
        f"D{d0}:D{d1}", FormulaRule(formula=[f"AND(ISNUMBER(D{d0}),D{d0}>=1)"], fill=fill("D9EAD3")))
    ws.conditional_formatting.add(
        f"D{R['h1']}:D{R[f'h{N_HABIT}']}",
        ColorScaleRule(start_type="num", start_value=0, start_color="F4CCCC",
                       mid_type="num", mid_value=0.5, mid_color="FFF2CC",
                       end_type="num", end_value=1, end_color="B7E1CD"))
    widths(ws, {"A": 28, "B": 11, "C": 11, "D": 8})
    for col in day_cols:
        ws.column_dimensions[col].width = 12
    ws.freeze_panes = "E6"
    ws.page_setup.orientation = "landscape"

    # ============================================================ 経費ログ
    ws = ws_exp
    title(ws, "A1", "経費ログ（費用を1件1行で入力）", "G1")
    header(ws, 2, ["日付", "勘定科目", "内容・支払先", "金額（円・税込）", "支払方法", "領収書", "メモ"],
           height=30)
    ex = ["記入例", "販管費：旅費交通費", "A社訪問（電車往復）", 1280, "電子マネー・QR", "○", ""]
    for i, v in enumerate(ex):
        put(ws, f"{get_column_letter(i + 1)}3", v, color="7F7F7F", italic=True, bg=C_EXAMPLE,
            border=True, fmt=FMT_YEN if i == 3 else None)
    list_validation(ws, ACC_NAMES, ["B3:B2000"])
    list_validation(ws, PAY_LIST, ["E3:E2000"])
    list_validation(ws, '"○,×"', ["F3:F2000"])
    nonneg_validation(ws, ["D3:D2000"])
    ws.column_dimensions["A"].number_format = FMT_DATE
    ws.column_dimensions["D"].number_format = FMT_YEN
    note(ws, "I2", f"日付・勘定科目・金額は必須。科目が「{COGS_PREFIX}」で始まる費用は売上原価、それ以外は販管費になる。"
         "本日の日報の売上原価・販管費と、月次集計に自動で反映される。「記入例」の行は集計されない。", "I6")
    widths(ws, {"A": 12, "B": 20, "C": 30, "D": 14, "E": 16, "F": 8, "G": 30, "H": 3, "I": 40})
    ws.freeze_panes = "A3"
    ws.sheet_properties.tabColor = "F1C232"

    # ============================================================ 案件パイプライン
    ws = ws_deal
    title(ws, "A1", "案件パイプライン（案件ごとに1行。フェーズと次のアクションを毎日更新）", "K1")
    header(ws, 2, ["登録日", "顧客名", "案件名・内容", "フェーズ", "見込金額（円）", "受注予定日",
                   "次のアクション", "期日", "担当", "流入経路", "メモ"], height=30)
    ex = ["記入例", "株式会社サンプル", "業務システム導入", "提案済", 1500000, date(2026, 11, 30),
          "見積を再提出する", date(2026, 10, 5), "自分", "既存顧客の紹介", ""]
    for i, v in enumerate(ex):
        put(ws, f"{get_column_letter(i + 1)}3", v, color="7F7F7F", italic=True, bg=C_EXAMPLE,
            border=True, fmt=FMT_YEN if i == 4 else (FMT_DATE if i in (5, 7) else None))
    list_validation(ws, PHASE_LIST, ["D3:D2000"])
    nonneg_validation(ws, ["E3:E2000"])
    for col in "AFH":
        ws.column_dimensions[col].number_format = FMT_DATE
    ws.column_dimensions["E"].number_format = FMT_YEN
    ws.conditional_formatting.add(
        "H3:H2000",
        FormulaRule(formula=['AND(ISNUMBER($A3),$H3<>"",$H3<TODAY(),$D3<>"受注",$D3<>"失注")'],
                    font=Font(color="C00000", bold=True)))
    # フェーズ別集計
    put(ws, "M2", "フェーズ別の集計", bold=True, color=NAVY)
    header(ws, 3, ["フェーズ", "件数", "見込金額", "確度", "加重金額"], start_col=13)
    for i in range(len(PHASES)):
        r = 4 + i
        calc(ws, f"M{r}", f"={q(S_SET)}!K{5 + i}&\"\"", link=True)
        calc(ws, f"N{r}", f'=COUNTIFS($A:$A,">0",$D:$D,M{r})', fmt=FMT_CNT)
        calc(ws, f"O{r}", f'=SUMIFS($E:$E,$A:$A,">0",$D:$D,M{r})', fmt=FMT_YEN)
        calc(ws, f"P{r}", f"={q(S_SET)}!L{5 + i}", fmt="0%", link=True)
        calc(ws, f"Q{r}", f"=O{r}*P{r}", fmt=FMT_YEN)
    rt = 4 + len(PHASES)
    put(ws, f"M{rt}", "合計", bold=True, border=True)
    calc(ws, f"N{rt}", f"=SUM(N4:N{rt - 1})", fmt=FMT_CNT, bold=True)
    calc(ws, f"O{rt}", f"=SUM(O4:O{rt - 1})", fmt=FMT_YEN, bold=True)
    put(ws, f"P{rt}", None, border=True)
    calc(ws, f"Q{rt}", f"=SUM(Q4:Q{rt - 1})", fmt=FMT_YEN, bold=True)
    put(ws, f"M{rt + 2}", "期日を過ぎた案件", border=True)
    calc(ws, f"N{rt + 2}",
         f'=COUNTIFS($A:$A,">0",$H:$H,"<"&TODAY(),$D:$D,"<>"&{PH_WON},$D:$D,"<>"&{PH_LOST})',
         fmt="0", bold=True)
    note(ws, f"O{rt + 2}", "受注・失注以外で、期日が今日より前の案件数", f"Q{rt + 2}", 28)
    note(ws, f"M{rt + 4}", "加重金額＝見込金額×確度（確度は設定シート）。「記入例」の行は集計されない。",
         f"Q{rt + 4}", 28)
    widths(ws, {"A": 12, "B": 18, "C": 24, "D": 11, "E": 14, "F": 12, "G": 24, "H": 12, "I": 9,
                "J": 14, "K": 22, "L": 3, "M": 14, "N": 7, "O": 14, "P": 7, "Q": 14})
    ws.freeze_panes = "A3"
    ws.sheet_properties.tabColor = "F1C232"

    # ============================================================ 10年計画
    ws = ws_plan
    title(ws, "A1", "10年計画（年商1000億円からの逆算）", "N1")
    note(ws, "A2", "設定シートの「1年目の売上目標」と「10年目の売上目標」から、毎年同じ成長率で伸ばした場合の目標を計算する。"
         "実績は日次ログの売上高から自動で集計する。", "N2", 18)
    put(ws, "A4", "1年目の売上目標", border=True)
    calc(ws, "B4", f'={SET["t1"]}', fmt=FMT_DEC, link=True)
    put(ws, "C4", "億円", border=True, align="center")
    put(ws, "A5", "10年目の売上目標", border=True)
    calc(ws, "B5", f'={SET["t10"]}', fmt=FMT_DEC, link=True)
    put(ws, "C5", "億円", border=True, align="center")
    put(ws, "A6", "必要な成長率（毎年）", border=True, bold=True)
    calc(ws, "B6", "=IF(B4<=0,0,(B5/B4)^(1/9)-1)", fmt=FMT_PCT, bold=True)
    put(ws, "C6", "％", border=True, align="center")
    note(ws, "D6", "この率で毎年伸ばすと、10年目に目標に届く（1年目から10年目までの9回分の成長）", "N6")
    put(ws, "A7", "10年目の1日あたり売上", border=True)
    calc(ws, "B7", "=B5/365", fmt="#,##0.00")
    put(ws, "C7", "億円", border=True, align="center")
    header(ws, 9, ["年目", "期間開始", "期間終了", "売上目標（億円）", "月平均（億円）", "1日あたり（万円）",
                   "実績（億円）", "達成率", "必要な受注件数（年）", "必要な商談数（年）",
                   "必要な新規アプローチ数（年）", "1日あたり必要なアプローチ数", "必要な営業担当者数（目安）",
                   "重点テーマ（例・書き換えてよい）"], height=45)
    for n in range(1, 11):
        r = 9 + n
        put(ws, f"A{r}", n, border=True, align="center", fmt='0"年目"')
        calc(ws, f"B{r}", f'=EDATE({SET["start"]},12*(A{r}-1))', fmt=FMT_DATE, link=True)
        calc(ws, f"C{r}", f'=EDATE({SET["start"]},12*A{r})-1', fmt=FMT_DATE, link=True)
        calc(ws, f"D{r}", f"=$B$4*(1+$B$6)^(A{r}-1)", fmt=FMT_DEC)
        calc(ws, f"E{r}", f"=D{r}/12", fmt="#,##0.00")
        calc(ws, f"F{r}", f"=D{r}*10^4/365", fmt=FMT_YEN)
        calc(ws, f"G{r}", f'=SUMIFS({log_sales},{log_a},">="&B{r},{log_a},"<="&C{r})/10^8',
             fmt="#,##0.00;▲#,##0.00;\"-\"", link=True)
        calc(ws, f"H{r}", f"=IF(D{r}=0,0,G{r}/D{r})", fmt=FMT_PCT)
        calc(ws, f"I{r}", f'=IF({SET["price"]}=0,0,D{r}*10^8/{SET["price"]})', fmt=FMT_CNT, link=True)
        calc(ws, f"J{r}", f'=IF({SET["close"]}=0,0,I{r}/{SET["close"]})', fmt=FMT_CNT, link=True)
        calc(ws, f"K{r}", f'=IF({SET["meet"]}=0,0,J{r}/{SET["meet"]})', fmt=FMT_CNT, link=True)
        calc(ws, f"L{r}", f"=K{r}/365", fmt=FMT_DEC)
        calc(ws, f"M{r}", f'=IF({SET["cap"]}=0,0,L{r}/{SET["cap"]})', fmt=FMT_DEC, link=True)
        inp(ws, f"N{r}", THEMES[n - 1], wrap=True)
        ws.row_dimensions[r].height = 30
    ws.conditional_formatting.add(
        "H10:H19", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=fill("D9EAD3")))
    note(ws, "A21", "必要な件数は、設定シートの平均受注単価・受注率・商談化率から計算する。単価や転換率が上がると、必要な件数と人数は減る。"
         "年商が大きくなるほど、1人の行動量ではなく、組織・単価・販路で売上をつくる必要があることがわかる。", "N21", 30)
    widths(ws, {"A": 22, "B": 12, "C": 12, "D": 11, "E": 10, "F": 11, "G": 10, "H": 9, "I": 11,
                "J": 11, "K": 13, "L": 12, "M": 12, "N": 60})
    ws.freeze_panes = "B10"

    # ============================================================ 月次集計
    ws = ws_month
    title(ws, "A1", "月次集計（日次ログ・経費ログから自動集計）", "S1")
    note(ws, "A2", "表示する計画年目を変えると、その年の12か月を表示する。売上・営業実績は日次ログ、原価・販管費は経費ログから集計。",
         "S2", 18)
    put(ws, "A4", "表示する計画年目", border=True, bold=True)
    inp(ws, "B4", 1, fmt='0"年目"', align="center")
    list_validation(ws, '"1,2,3,4,5,6,7,8,9,10"', ["B4"])
    put(ws, "A5", "期間", border=True)
    calc(ws, "B5", f'=EDATE({SET["start"]},12*($B$4-1))', fmt=FMT_DATE, link=True)
    put(ws, "C5", "〜", align="center")
    calc(ws, "D5", f'=EDATE({SET["start"]},12*$B$4)-1', fmt=FMT_DATE, link=True)
    month_heads = ["月", "売上高", "売上目標", "達成率", "売上原価", "粗利", "粗利率", "販管費",
                   "営業利益", "営業利益率", "新規アプローチ", "商談数", "受注数", "受注金額", "商談化率",
                   "受注率", "平均受注単価", "日報提出日数", "日課達成率（平均）"]
    header(ws, 7, month_heads, height=32)

    def lsum(key, r):
        return (f'SUMIFS({q(S_LOG)}!${L[key]}:${L[key]},{log_a},">="&$A{r},'
                f'{log_a},"<="&EOMONTH($A{r},0))')

    for k in range(12):
        r = 8 + k
        calc(ws, f"A{r}", "=DATE(YEAR($B$5),MONTH($B$5),1)" if k == 0 else f"=EDATE(A{r - 1},1)",
             fmt=FMT_MONTH, align="center")
        exp_rng = f'{exp_a},">="&$A{r},{exp_a},"<="&EOMONTH($A{r},0)'
        f = {
            "B": f"={lsum('sales', r)}",
            "C": f"=INDEX({plan_d},$B$4)*10^8/12",
            "D": f"=IF(C{r}=0,0,B{r}/C{r})",
            "E": f'=SUMIFS({exp_d},{exp_rng},{exp_b},"{COGS_PREFIX}*")',
            "F": f"=B{r}-E{r}",
            "G": f"=IF(B{r}=0,0,F{r}/B{r})",
            "H": f"=SUMIFS({exp_d},{exp_rng})-E{r}",
            "I": f"=F{r}-H{r}",
            "J": f"=IF(B{r}=0,0,I{r}/B{r})",
            "K": f"={lsum('approach', r)}",
            "L": f"={lsum('meeting', r)}",
            "M": f"={lsum('won', r)}",
            "N": f"={lsum('won_amt', r)}",
            "O": f"=IF(K{r}=0,0,L{r}/K{r})",
            "P": f"=IF(L{r}=0,0,M{r}/L{r})",
            "Q": f"=IF(M{r}=0,0,N{r}/M{r})",
            "R": f'=COUNTIFS({log_a},">="&$A{r},{log_a},"<="&EOMONTH($A{r},0))',
            "S": f"=IF(R{r}=0,0,{lsum('habit_rate', r)}/R{r})",
        }
        for col, formula in f.items():
            pct = col in "DGJOPS"
            calc(ws, f"{col}{r}", formula, fmt=FMT_PCT if pct else FMT_YEN,
                 link=col in "BCEHKLMNRS")
    rt = 20
    put(ws, f"A{rt}", "年間合計", bold=True, border=True, bg=C_SUB, align="center")
    for col in "BCEFHIKLMNR":
        calc(ws, f"{col}{rt}", f"=SUM({col}8:{col}19)", fmt=FMT_YEN, bold=True)
    ratio = {"D": "=IF(C20=0,0,B20/C20)", "G": "=IF(B20=0,0,F20/B20)",
             "J": "=IF(B20=0,0,I20/B20)", "O": "=IF(K20=0,0,L20/K20)",
             "P": "=IF(L20=0,0,M20/L20)", "S": "=IF(R20=0,0,SUMPRODUCT(S8:S19,R8:R19)/R20)"}
    for col, formula in ratio.items():
        calc(ws, f"{col}{rt}", formula, fmt=FMT_PCT, bold=True)
    calc(ws, "Q20", "=IF(M20=0,0,N20/M20)", fmt=FMT_YEN, bold=True)
    for col in "ABCDEFGHIJKLMNOPQRS":
        ws[f"{col}{rt}"].fill = fill(C_SUB)
    ws.conditional_formatting.add(
        "D8:D20", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=fill("D9EAD3")))
    ws.conditional_formatting.add(
        "I8:I20", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000")))

    # 日課の実行率（月別）
    g0 = 22
    section(ws, g0, "日課の実行率（○をつけた日数÷日報を提出した日数）", "A", "N")
    header(ws, g0 + 1, ["No", "日課"])
    for k in range(12):
        col = get_column_letter(3 + k)
        calc(ws, f"{col}{g0 + 1}",
             "=DATE(YEAR($B$5),MONTH($B$5),1)" if k == 0
             else f"=EDATE({get_column_letter(2 + k)}{g0 + 1},1)", fmt=FMT_MONTH, align="center")
        ws[f"{col}{g0 + 1}"].font = font(True, 10, NAVY)
        ws[f"{col}{g0 + 1}"].fill = fill(C_HEAD)
    put(ws, f"B{g0 + 2}", "日報を提出した日数", border=True, italic=True)
    put(ws, f"A{g0 + 2}", None, border=True)
    habit_cols = f"{q(S_LOG)}!${L['h1']}:${L[f'h{N_HABIT}']}"
    for k in range(12):
        col = get_column_letter(3 + k)
        calc(ws, f"{col}{g0 + 2}",
             f'=COUNTIFS({log_a},">="&{col}${g0 + 1},{log_a},"<="&EOMONTH({col}${g0 + 1},0))',
             fmt="0", link=True, align="center")
    for i in range(N_HABIT):
        r = g0 + 3 + i
        put(ws, f"A{r}", i + 1, border=True, align="center")
        calc(ws, f"B{r}", f'={habit_ref("C", i)}&""', link=True)
        for k in range(12):
            col = get_column_letter(3 + k)
            calc(ws, f"{col}{r}",
                 f'=IF({col}${g0 + 2}=0,"",COUNTIFS({log_a},">="&{col}${g0 + 1},'
                 f'{log_a},"<="&EOMONTH({col}${g0 + 1},0),INDEX({habit_cols},0,$A{r}),"○")'
                 f"/{col}${g0 + 2})", fmt="0%;;\"-\"", align="center")
    gr0, gr1 = g0 + 3, g0 + 2 + N_HABIT
    ws.conditional_formatting.add(
        f"C{gr0}:N{gr1}",
        ColorScaleRule(start_type="num", start_value=0, start_color="F4CCCC",
                       mid_type="num", mid_value=0.5, mid_color="FFF2CC",
                       end_type="num", end_value=1, end_color="B7E1CD"))
    widths(ws, {"A": 12, "B": 30})
    for k in range(17):
        ws.column_dimensions[get_column_letter(3 + k)].width = 11
    ws.freeze_panes = "B8"

    # ============================================================ 使い方
    ws = ws_howto
    title(ws, "A1", "日報・日次決算テンプレート（年商1000億円への10年計画）", "C1")
    note(ws, "A2", "毎日、日報・営業実績・日次決算・日課チェックを1枚に入力して提出し、日次ログにためる。"
         "月間日報・月次集計・10年計画は、ためたデータから自動で集計される。", "C2", 30)
    section(ws, 4, "毎日の流れ", "A", "C")
    flow = [
        ("1. 朝", "「本日の日報」で日付を確認し、前回決めた今日のTop3と売上目標を見る。「毎日やること」で今日の行動を決める。"),
        ("2. 日中", "費用が出たら「経費ログ」に1件1行で記録する。商談が動いたら「案件パイプライン」のフェーズと次のアクションを更新する。"),
        ("3. 夜", "「本日の日報」の薄い黄色のセルに、売上・営業実績・日課の結果・振り返り（KPT）・明日のTop3を入力する。売上原価・販管費・利益は経費ログから自動で計算される。"),
        ("4. 提出", "メニュー「日報」→「本日の日報を提出」を押す（Apps Scriptの設定が必要。下記3）。スクリプトを使わない場合は「提出データ」シートの2行目をコピーし、「日次ログ」の最終行の下に値のみ貼り付け（Ctrl+Shift+V）。"),
        ("5. 週に1回", "「月間日報」でその月の日報を1日1列で見返す。「月次集計」で売上・利益・日課の実行率を、「10年計画」で今年の目標との差を確認する。続かない日課は原因を1つ決めて手を打つ。"),
    ]
    for i, (a, b) in enumerate(flow):
        r = 5 + i
        put(ws, f"A{r}", a, bold=True, border=True, valign="top")
        put(ws, f"B{r}", b, border=True, wrap=True, valign="top")
        ws.merge_cells(f"B{r}:C{r}")
        ws.row_dimensions[r].height = 48 if a == "5. 週に1回" else 42
    section(ws, 11, "セルの色", "A", "C")
    inp(ws, "A12", "入力値")
    put(ws, "B12", "薄い黄色：入力するセル（青い文字が入力した値）", border=True)
    put(ws, "A13", "計算値", border=True)
    put(ws, "B13", "白：自動で計算されるセル（入力しない）", border=True)
    put(ws, "A14", "参照値", color="008000", border=True)
    put(ws, "B14", "緑の文字：ほかのシートから参照している値", border=True)
    for r in (12, 13, 14):
        ws.merge_cells(f"B{r}:C{r}")
    section(ws, 16, "最初に1回だけやること", "A", "C")
    first = [
        ("1", "「設定」シートの薄い黄色のセル（計画開始日、1年目の売上目標、平均受注単価、受注率など）を自分の数字に書き換える。"),
        ("2", "「毎日やること」の項目と目安を、自分の事業に合わせて見直す（本日の日報・日次ログ・月次集計に自動で反映される）。"),
        ("3", "提出ボタンを使う場合は、メニュー「拡張機能」→「Apps Script」を開き、最初からあるコードを消して、提出スクリプト（Googleドライブの「日報 提出スクリプト（Apps Script用）」、またはGitHubの daily-report/gas/Code.gs）を貼り付けて保存する。シートを再読み込みするとメニューに「日報」が表示される。初回だけGoogleの承認画面が出るので、自分のアカウントで許可する。"),
        ("4", "提出時にメールも送る場合は「設定」の提出先メールアドレスを入力する。毎晩の未提出リマインドは、メニュー「日報」→「毎日のリマインドを設定」。"),
    ]
    for i, (a, b) in enumerate(first):
        r = 17 + i
        put(ws, f"A{r}", a, border=True, align="center", valign="top")
        put(ws, f"B{r}", b, border=True, wrap=True, valign="top")
        ws.merge_cells(f"B{r}:C{r}")
        ws.row_dimensions[r].height = 60 if a == "3" else 32
    section(ws, 22, "シートの一覧", "A", "C")
    header(ws, 23, ["シート", "役割", "いつ入力するか"])
    sheets = [
        (S_FORM, "毎日の入力画面。日報・営業実績・日次決算・日課チェック", "毎日"),
        (S_DAYS, "提出した日報を1か月分、1日1列で表示（月計・月目標・達成率つき）", "見るだけ（表示する月は変更可）"),
        (S_HABIT, "年商1000億に向けて毎日やることの一覧と理由", "必要に応じて見直す"),
        (S_EXP, "費用を1件1行で記録。日次決算と月次集計に反映", "費用が出たとき"),
        (S_DEAL, "案件ごとのフェーズ・見込金額・次のアクション", "商談が動いたとき"),
        (S_MONTH, "月ごとの売上・利益・営業数字・日課の実行率", "見るだけ（表示する年目は変更可）"),
        (S_PLAN, "1〜10年目の売上目標と実績、必要な件数と人数の逆算", "見るだけ（重点テーマは書き換え可）"),
        (S_LOG, "提出した日報が1日1行でたまる記録", "提出時に自動"),
        (S_SUBMIT, "日次ログに送る1行分のデータ（自動作成）", "触らない"),
        (S_SET, "目標・前提条件・選択肢のリスト", "最初に1回（毎月見直す）"),
    ]
    for i, (a, b, c) in enumerate(sheets):
        r = 24 + i
        put(ws, f"A{r}", a, border=True, bold=True)
        put(ws, f"B{r}", b, border=True, wrap=True)
        put(ws, f"C{r}", c, border=True, wrap=True)
    section(ws, 35, "補足", "A", "C")
    notes = [
        "日次ログ・経費ログ・案件パイプラインの「記入例」の行は集計されない。不要なら行ごと削除してよい。",
        "経費は経費ログが正本。提出後に経費を追加・修正しても月次集計には反映される（日次ログの費用欄は提出時点の記録）。",
        "目標の数字（仮の値）は設定シートで変更できる。変更すると10年計画・本日の日報・月次集計が連動する。",
        "同じ日付の日報をもう一度提出すると、上書きするか確認が出る。",
    ]
    for i, t in enumerate(notes):
        r = 36 + i
        put(ws, f"A{r}", "・", align="right", valign="top")
        put(ws, f"B{r}", t, wrap=True, valign="top")
        ws.merge_cells(f"B{r}:C{r}")
        ws.row_dimensions[r].height = 28
    widths(ws, {"A": 16, "B": 60, "C": 34})
    ws.sheet_properties.tabColor = NAVY

    wb.active = wb.sheetnames.index(S_FORM)
    for s in wb.worksheets:
        s.sheet_view.tabSelected = s.title == S_FORM
    return wb, L


# ---------------------------------------------------------------- ファイルを小さくする
# Google ドライブへのアップロードを軽くするため、保存後の xlsx を後処理する。
# ・同じ形の数式が並ぶ範囲を「共有数式」（Excel が行や列をコピーしたときと同じ形式）にまとめる
# ・シートごとの文字列を1つの共有文字列表（sharedStrings.xml）にまとめる
# ・数式の空のキャッシュ値 <v /> と、任意の部品（docProps、テーマ）を取り除く
CELL_RE = re.compile(r'<c r="([A-Z]+)(\d+)"([^>]*)><f>(.*?)</f>(?:<v ?/>|<v></v>)</c>')
CHARREF_RE = re.compile(r"&#(x[0-9a-fA-F]+|\d+);")


def _plain_text(xml):
    """日本語などの数値文字参照（&#26085; など）を UTF-8 の文字に戻す。"""
    def repl(m):
        code = m.group(1)
        n = int(code[1:], 16) if code.startswith("x") else int(code)
        return chr(n) if n > 127 else m.group(0)
    return CHARREF_RE.sub(repl, xml)


def _share_formulas(xml):
    cells = {}
    for m in CELL_RE.finditer(xml):
        col, row, attrs, f = m.groups()
        cells[(column_index_from_string(col), int(row))] = (attrs, html.unescape(f))

    def same_shape(anchor, target):
        (ac, ar), (tc, tr) = anchor, target
        if target not in cells:
            return False
        src = "=" + cells[anchor][1]
        moved = Translator(src, origin=f"{get_column_letter(ac)}{ar}").translate_formula(
            f"{get_column_letter(tc)}{tr}")
        return moved == "=" + cells[target][1]

    group_of, groups = {}, []
    for key in sorted(cells, key=lambda k: (k[1], k[0])):
        if key in group_of:
            continue
        c0, r0 = key
        r1 = r0
        while (c0, r1 + 1) not in group_of and same_shape(key, (c0, r1 + 1)):
            r1 += 1
        c1 = c0
        while all((c1 + 1, r) not in group_of and same_shape(key, (c1 + 1, r))
                  for r in range(r0, r1 + 1)):
            c1 += 1
        if (r1 - r0 + 1) * (c1 - c0 + 1) < 2:
            continue
        si = len(groups)
        groups.append((key, f"{get_column_letter(c0)}{r0}:{get_column_letter(c1)}{r1}"))
        for c in range(c0, c1 + 1):
            for r in range(r0, r1 + 1):
                group_of[(c, r)] = si

    def repl(m):
        col, row, attrs, f = m.groups()
        key = (column_index_from_string(col), int(row))
        if key not in group_of:
            return f'<c r="{col}{row}"{attrs}><f>{f}</f></c>'
        si = group_of[key]
        anchor, ref = groups[si]
        if key == anchor:
            return f'<c r="{col}{row}"{attrs}><f t="shared" ref="{ref}" si="{si}">{f}</f></c>'
        return f'<c r="{col}{row}"{attrs}><f t="shared" si="{si}"/></c>'

    return CELL_RE.sub(repl, xml)


INLINE_RE = re.compile(r'<c r="([A-Z]+\d+)"((?: s="\d+")?) t="inlineStr"><is><t([^>]*)>(.*?)</t></is></c>')
SS_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def compact_xlsx(path):
    src = zipfile.ZipFile(path)
    parts = {i.filename: src.read(i.filename) for i in src.infolist()}
    src.close()
    for name in ("docProps/app.xml", "docProps/core.xml", "xl/theme/theme1.xml"):
        parts.pop(name, None)
    parts["_rels/.rels"] = re.sub(rb'<Relationship [^>]*docProps[^>]*/>', b"", parts["_rels/.rels"])
    ct = re.sub(rb'<Override PartName="/(docProps|xl/theme)/[^>]*/>', b"", parts["[Content_Types].xml"])
    rels = re.sub(rb'<Relationship [^>]*theme1.xml[^>]*/>', b"", parts["xl/_rels/workbook.xml.rels"])
    styles = parts["xl/styles.xml"].replace(b'<color theme="1" />', b'<color rgb="FF000000" />')
    parts["xl/styles.xml"] = styles.replace(b'<scheme val="minor" />', b"")

    strings, index = [], {}

    def to_shared(m):
        ref, style, t_attrs, text = m.groups()
        key = (t_attrs, text)
        if key not in index:
            index[key] = len(strings)
            strings.append(key)
        return f'<c r="{ref}"{style} t="s"><v>{index[key]}</v></c>'

    for name in sorted(parts):
        if name.startswith("xl/worksheets/sheet"):
            xml = _share_formulas(_plain_text(parts[name].decode("utf-8")))
            parts[name] = INLINE_RE.sub(to_shared, xml).encode("utf-8")
    if strings:
        items = "".join(f"<si><t{attrs}>{text}</t></si>" for attrs, text in strings)
        parts["xl/sharedStrings.xml"] = (
            f'<sst xmlns="{SS_NS}" uniqueCount="{len(strings)}">{items}</sst>').encode("utf-8")
        ct = ct.replace(b"</Types>", b'<Override PartName="/xl/sharedStrings.xml" ContentType="'
                        b'application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml" />'
                        b"</Types>")
        rels = rels.replace(b"</Relationships>", f'<Relationship Type="{REL_NS}/sharedStrings" '
                            f'Target="sharedStrings.xml" Id="rIdStrings" /></Relationships>'.encode())
    parts["[Content_Types].xml"] = ct
    parts["xl/_rels/workbook.xml.rels"] = rels
    # 標準的な順番（[Content_Types].xml が先頭）で、日時を固定して書く。同じ内容なら毎回同じファイルになる。
    with zipfile.ZipFile(path, "w") as out:
        for name in sorted(parts, key=_part_order):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            out.writestr(info, parts[name], compresslevel=9)


PART_ORDER = ("[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels",
              "xl/styles.xml", "xl/sharedStrings.xml")


def _part_order(name):
    if name in PART_ORDER:
        return (PART_ORDER.index(name), 0)
    m = re.search(r"(\d+)\.xml$", name)
    return (len(PART_ORDER), int(m.group(1)) if m else 0)


if __name__ == "__main__":
    wb, _ = build()
    wb.save(OUT)
    compact_xlsx(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size:,} bytes)")
